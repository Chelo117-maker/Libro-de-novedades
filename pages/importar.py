# Importación de guardias desde archivos externos (.txt, .pdf, .xlsx, .docx).

from nicegui import ui
from operario import Operario
from guardia import Guardia
from components.sidebar import crear_sidebar, crear_header, ESTADO_CONFIG
from components.ayuda import mostrar_dialogo_ayuda
import novedades


def configurar_tema() -> None:
    """Activa el modo oscuro y el color primario."""
    ui.dark_mode().enable()
    ui.colors(primary="#1976D2")


def parsear_txt(contenido: str) -> dict:
    """
    Lee el contenido de un archivo y extrae los datos de la guardia.
    Soporta las secciones: FECHA, FRANJA, OPERARIOS, HERRAMIENTAS, NOVEDADES.
    """
    datos: dict = {
        "fecha": "",
        "franja": "",
        "operarios": [],
        "herramientas": [],
        "novedades": [],
    }
    seccion = None
    for linea in contenido.splitlines():
        linea = linea.strip()
        if not linea:
            continue
        if linea.startswith("FECHA:"):
            datos["fecha"] = linea.replace("FECHA:", "").strip()
        elif linea.startswith("FRANJA:"):
            datos["franja"] = linea.replace("FRANJA:", "").strip()
        elif linea == "OPERARIOS:":
            seccion = "operarios"
        elif linea == "HERRAMIENTAS:":
            seccion = "herramientas"
        elif linea == "NOVEDADES:":
            seccion = "novedades"
        elif seccion == "operarios" and " - " in linea:
            partes = linea.split(" - ", 1)
            nc = partes[1].strip().split(" ", 1)
            datos["operarios"].append(
                {
                    "legajo": partes[0].strip(),
                    "nombre": nc[0],
                    "apellido": nc[1] if len(nc) > 1 else "",
                }
            )
        elif seccion == "herramientas" and " - " in linea:
            partes = linea.split(" - ", 1)
            datos["herramientas"].append(
                {
                    "nombre": partes[0].strip(),
                    "estado": partes[1].strip(),
                }
            )
        elif seccion == "novedades":
            datos["novedades"].append(linea)
    return datos


@ui.page("/importar")
def pagina_importar() -> None:
    """Pantalla para importar guardias desde archivos externos."""
    configurar_tema()
    drawer = crear_sidebar("importar")
    crear_header(
        drawer, "Importar novedades", mostrar_ayuda=True, fn_ayuda=mostrar_dialogo_ayuda
    )

    estado = {"datos": None}

    with ui.column().classes("q-pa-md full-width"):

        # ── Sección de subida ─────────────────────────────────────────────────
        ui.label("SUBIR ARCHIVO DE NOVEDADES").classes("text-caption text-grey")
        with ui.card().classes("full-width q-mb-md"):
            ui.label(
                "Formatos soportados: .txt (recomendado), .pdf, .xlsx y .docx"
            ).classes("text-caption text-grey q-mb-sm")
            ui.label(
                "Descargá la plantilla desde Ayuda → Formato antes de completar el archivo."
            ).classes("text-caption text-grey q-mb-md")

            preview = ui.column().classes("full-width")

            def mostrar_preview(datos: dict) -> None:
                """Muestra una vista previa de los datos parseados."""
                estado["datos"] = datos
                preview.clear()
                with preview:
                    if not datos["fecha"] or not datos["franja"]:
                        ui.label(
                            "⚠️ No se encontró fecha o franja válida. "
                            "Verificá que el formato sea correcto."
                        ).classes("text-warning text-body2 q-mt-sm")
                        return

                    ui.label("VISTA PREVIA").classes(
                        "text-caption text-grey q-mt-md q-mb-sm"
                    )

                    with ui.card().classes("full-width q-mb-sm"):
                        ui.label(f"📅 Fecha: {datos['fecha']}").classes("text-body2")
                        ui.label(f"🕐 Franja: {datos['franja']}").classes("text-body2")

                    with ui.card().classes("full-width q-mb-sm"):
                        ui.label("Operarios:").classes("text-caption text-grey")
                        if datos["operarios"]:
                            for op in datos["operarios"]:
                                ui.label(
                                    f"  [{op['legajo']}] {op['nombre']} {op['apellido']}"
                                ).classes("text-body2")
                        else:
                            ui.label("  Sin operarios detectados.").classes(
                                "text-grey text-caption"
                            )

                    with ui.card().classes("full-width q-mb-sm"):
                        ui.label("Novedades:").classes("text-caption text-grey")
                        if datos["novedades"]:
                            for nov in datos["novedades"]:
                                ui.label(f"  • {nov}").classes("text-body2")
                        else:
                            ui.label("  Sin novedades detectadas.").classes(
                                "text-grey text-caption"
                            )

                    if datos["herramientas"]:
                        with ui.card().classes("full-width q-mb-sm"):
                            ui.label("Herramientas:").classes("text-caption text-grey")
                            for h in datos["herramientas"]:
                                cfg = ESTADO_CONFIG.get(
                                    h["estado"], ESTADO_CONFIG["Operativa"]
                                )
                                with ui.row().classes("items-center q-gutter-sm"):
                                    ui.icon(cfg["icon"]).classes(f"text-{cfg['color']}")
                                    ui.label(f"{h['nombre']} — {h['estado']}").classes(
                                        "text-body2"
                                    )

                    ui.button(
                        "✓ Confirmar e importar", on_click=confirmar_importacion
                    ).props("color=positive rounded").classes("full-width q-mt-sm")

            def confirmar_importacion() -> None:
                """Registra la guardia importada en el sistema."""
                datos = estado["datos"]
                if not datos:
                    ui.notify("No hay datos para importar", color="negative")
                    return
                if not datos.get("fecha") or not datos.get("franja"):
                    ui.notify(
                        "El archivo no tiene fecha o franja válida", color="negative"
                    )
                    return
                g = Guardia(datos["fecha"], datos["franja"])
                for op in datos["operarios"]:
                    g.agregar_operario(
                        Operario(op["legajo"], op["nombre"], op["apellido"])
                    )
                for nov in datos["novedades"]:
                    g.agregar_novedad(nov)
                novedades.agregar_guardia(g)
                for h in datos.get("herramientas", []):
                    novedades.agregar_herramienta(h["nombre"], h["estado"])
                ui.notify("Guardia importada correctamente ✓", color="positive")
                preview.clear()
                estado["datos"] = None
                ui.navigate.to("/")

            async def procesar_archivo(e) -> None:
                """Lee el archivo subido y lo parsea según su extensión."""
                try:
                    archivo = e.file
                    nombre = (
                        getattr(archivo, "filename", None)
                        or getattr(archivo, "name", None)
                        or "archivo.txt"
                    )
                    nombre = nombre.lower()
                    datos_b = await archivo.read()

                    contenido = ""

                    if nombre.endswith(".txt"):
                        contenido = datos_b.decode("utf-8", errors="ignore")

                    elif nombre.endswith(".pdf"):
                        try:
                            import io
                            import pdfplumber

                            with pdfplumber.open(io.BytesIO(datos_b)) as pdf:
                                contenido = "\n".join(
                                    p.extract_text() or "" for p in pdf.pages
                                )
                        except ImportError:
                            ui.notify("Falta instalar pdfplumber", color="warning")
                            return

                    elif nombre.endswith(".xlsx"):
                        try:
                            import io
                            import openpyxl

                            wb = openpyxl.load_workbook(io.BytesIO(datos_b))
                            ws = wb.active
                            contenido = "\n".join(
                                " - ".join(str(c.value or "") for c in row)
                                for row in ws.iter_rows()
                            )
                        except ImportError:
                            ui.notify("Falta instalar openpyxl", color="warning")
                            return

                    elif nombre.endswith(".docx"):
                        try:
                            import io
                            import docx

                            doc = docx.Document(io.BytesIO(datos_b))
                            contenido = "\n".join(p.text for p in doc.paragraphs)
                        except ImportError:
                            ui.notify("Falta instalar python-docx", color="warning")
                            return

                    else:
                        ui.notify(
                            "Formato no soportado. Usá .txt, .pdf, .xlsx o .docx",
                            color="negative",
                        )
                        return

                    if contenido.strip():
                        datos = parsear_txt(contenido)
                        mostrar_preview(datos)
                        ui.notify("Archivo leído correctamente ✓", color="positive")
                    else:
                        ui.notify(
                            "No se pudo leer contenido del archivo", color="negative"
                        )

                except Exception as ex:
                    import traceback

                    print(traceback.format_exc())
                    ui.notify(f"Error: {str(ex)}", color="negative")

        ui.upload(
            label="Seleccionar archivo",
            on_upload=procesar_archivo,
            auto_upload=True,
            max_file_size=5_000_000,
        ).props('accept=".txt,.pdf,.xlsx,.docx"').classes("full-width")
